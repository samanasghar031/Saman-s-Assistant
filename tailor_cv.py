import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document


BASE_FOLDER = Path(__file__).resolve().parent
MASTER_CV = BASE_FOLDER / "Master_CV.docx"
VERIFICATION_FILE = BASE_FOLDER / "verified_professor.json"
EVIDENCE_FILE = BASE_FOLDER / "student_research_evidence.json"
OUTPUT_FOLDER = BASE_FOLDER / "tailored_cvs"

MANIFEST_FILE = (
    OUTPUT_FOLDER
    / "latest_tailored_cv_manifest.json"
)

MIN_PROFILE_WORDS = 45
MAX_PROFILE_WORDS = 90

FORBIDDEN_PROFILE_TERMS = (
    "applicant",
    "doctoral",
    "funding",
    "phd",
    "professor",
    "recruiting",
    "recruitment",
    "supervision",
    "supervisor",
    "tailored",
    "targeted",
    "university",
)


def sha256_file(path):
    digest = hashlib.sha256()

    with path.open("rb") as file:
        for block in iter(
            lambda: file.read(1024 * 1024),
            b"",
        ):
            digest.update(block)

    return digest.hexdigest()


def load_json(path, description):
    if not path.exists():
        raise FileNotFoundError(
            f"{description} was not found: "
            f"{path.name}"
        )

    try:
        return json.loads(
            path.read_text(
                encoding="utf-8"
            )
        )

    except json.JSONDecodeError as error:
        raise ValueError(
            f"{description} is not valid JSON."
        ) from error


def safe_filename(value):
    value = re.sub(
        r"[^A-Za-z0-9_-]+",
        "_",
        value or "",
    )

    return (
        value.strip("_")
        or "Professor"
    )


def require_approved_inputs():
    verification = load_json(
        VERIFICATION_FILE,
        "verified_professor.json",
    )

    evidence = load_json(
        EVIDENCE_FILE,
        "student_research_evidence.json",
    )

    if (
        verification.get(
            "schema_version",
            0,
        )
        < 3
    ):
        raise ValueError(
            "The professor report came from "
            "the old verification workflow."
        )

    if (
        verification.get(
            "ready_for_contact"
        )
        != "YES"
    ):
        blockers = verification.get(
            "blockers",
            [],
        )

        details = (
            "; ".join(blockers)
            if blockers
            else "Verification failed."
        )

        raise ValueError(
            "This professor is not approved "
            "for contact. "
            + details
        )

    if (
        verification.get(
            "recruitment_reverified"
        )
        is not True
    ):
        raise ValueError(
            "The professor's 2027 recruitment "
            "evidence was not reverified."
        )

    if (
        verification.get(
            "research_fit_passed"
        )
        is not True
    ):
        raise ValueError(
            "The professor did not pass the "
            "strict research-fit check."
        )

    if not verification.get(
        "verified_publications"
    ):
        raise ValueError(
            "No professor-authored publication "
            "was verified."
        )

    for field, description in (
        (
            "candidate_id",
            "candidate ID",
        ),
        (
            "name",
            "professor name",
        ),
        (
            "research_topic",
            "verified research topic",
        ),
    ):
        if not str(
            verification.get(
                field,
                "",
            )
        ).strip():
            raise ValueError(
                f"The {description} is missing."
            )

    if not verification.get(
        "verified_overlap_terms"
    ):
        raise ValueError(
            "No verified research-overlap "
            "terms were recorded."
        )

    if (
        evidence.get(
            "approved_by_user"
        )
        is not True
    ):
        raise ValueError(
            "The student research evidence "
            "is not approved."
        )

    applicant_name = str(
        evidence.get(
            "applicant_name",
            "",
        )
    ).strip()

    if not applicant_name:
        raise ValueError(
            "Applicant name is missing from "
            "the approved evidence."
        )

    if not MASTER_CV.exists():
        raise FileNotFoundError(
            "Master_CV.docx was not found."
        )

    return verification, evidence


def locate_profile_paragraph(document):
    paragraphs = document.paragraphs

    heading_indexes = [
        index
        for index, paragraph
        in enumerate(paragraphs)
        if (
            paragraph.text.strip().upper()
            == "PROFILE"
        )
    ]

    if len(heading_indexes) != 1:
        raise ValueError(
            "Master_CV.docx must contain "
            "exactly one PROFILE heading."
        )

    for paragraph in paragraphs[
        heading_indexes[0] + 1:
    ]:
        if paragraph.text.strip():
            return paragraph

    raise ValueError(
        "No profile paragraph was found "
        "after PROFILE."
    )


def split_sentences(text):
    return [
        sentence.strip()
        for sentence in re.split(
            r"(?<=[.!?])\s+",
            text.strip(),
        )
        if sentence.strip()
    ]


def safe_research_interest(
    verification,
):
    topic = re.sub(
        r"\s+",
        " ",
        str(
            verification.get(
                "research_topic",
                "",
            )
        ),
    ).strip(
        " .,:;-"
    )

    if not topic:
        raise ValueError(
            "The verified research topic "
            "is missing."
        )

    if len(topic) > 120:
        raise ValueError(
            "The verified research topic "
            "is unexpectedly long."
        )

    if re.search(
        r"https?://|[.!?]",
        topic,
        re.IGNORECASE,
    ):
        raise ValueError(
            "The verified research topic "
            "contains unsafe text."
        )

    lowered = topic.casefold()

    for term in FORBIDDEN_PROFILE_TERMS:
        if re.search(
            rf"\b{re.escape(term)}\b",
            lowered,
        ):
            raise ValueError(
                "The verified research topic "
                f'contains "{term}", which is '
                "not allowed in the CV profile."
            )

    topic_words = set(
        re.findall(
            r"[a-z][a-z0-9-]+",
            lowered,
        )
    )

    overlap_terms = [
        str(value).casefold().strip()
        for value in verification.get(
            "verified_overlap_terms",
            [],
        )
        if str(value).strip()
    ]

    supported = False

    for overlap in overlap_terms:
        overlap_words = set(
            re.findall(
                r"[a-z][a-z0-9-]+",
                overlap,
            )
        )

        if (
            overlap in lowered
            or topic_words.intersection(
                overlap_words
            )
        ):
            supported = True
            break

    if not supported:
        raise ValueError(
            "The research topic is not "
            "supported by the verified "
            "research-overlap terms."
        )

    return topic


def build_simple_profile(
    original_profile,
    verification,
):
    sentences = split_sentences(
        original_profile
    )

    if len(sentences) < 4:
        raise ValueError(
            "The master profile must contain "
            "at least four sentences."
        )

    unchanged_opening = " ".join(
        sentences[:3]
    )

    topic = safe_research_interest(
        verification
    )

    return (
        unchanged_opening
        + " Interested in developing "
        "expertise in "
        + topic
        + "."
    )


def validate_profile(
    original_profile,
    proposed_profile,
):
    original_sentences = split_sentences(
        original_profile
    )

    proposed_sentences = split_sentences(
        proposed_profile
    )

    if len(proposed_sentences) != 4:
        raise ValueError(
            "The proposed profile must contain "
            "exactly four sentences."
        )

    if (
        proposed_sentences[:3]
        != original_sentences[:3]
    ):
        raise ValueError(
            "The first three master-profile "
            "sentences were changed."
        )

    if (
        "\n" in proposed_profile
        or "**" in proposed_profile
    ):
        raise ValueError(
            "Headings, markdown, and line "
            "breaks are not allowed."
        )

    if not proposed_profile.startswith(
        "Forensic Science researcher"
    ):
        raise ValueError(
            "The profile must begin with "
            '"Forensic Science researcher".'
        )

    if re.search(
        r"\b(?:I|me|my|mine)\b",
        proposed_profile,
    ):
        raise ValueError(
            "First-person language is not "
            "allowed in the profile."
        )

    lowered = proposed_profile.casefold()

    for term in FORBIDDEN_PROFILE_TERMS:
        if re.search(
            rf"\b{re.escape(term)}\b",
            lowered,
        ):
            raise ValueError(
                "The proposed profile contains "
                f'the forbidden term "{term}".'
            )

    word_count = len(
        proposed_profile.split()
    )

    if not (
        MIN_PROFILE_WORDS
        <= word_count
        <= MAX_PROFILE_WORDS
    ):
        raise ValueError(
            "The proposed profile has "
            f"{word_count} words; the allowed "
            f"range is {MIN_PROFILE_WORDS}-"
            f"{MAX_PROFILE_WORDS}."
        )


def replace_paragraph_text(
    paragraph,
    new_text,
):
    if paragraph.runs:
        paragraph.runs[0].text = (
            new_text
        )

        for run in paragraph.runs[1:]:
            run.text = ""

    else:
        paragraph.add_run(
            new_text
        )


def unique_output_path(
    applicant_name,
    professor_name,
):
    OUTPUT_FOLDER.mkdir(
        exist_ok=True
    )

    base_name = (
        f"{safe_filename(applicant_name)}_CV_"
        f"{safe_filename(professor_name)}"
    )

    output_path = (
        OUTPUT_FOLDER
        / f"{base_name}.docx"
    )

    if not output_path.exists():
        return output_path

    stamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    return (
        OUTPUT_FOLDER
        / f"{base_name}_{stamp}.docx"
    )


def verify_saved_copy(
    path,
    expected_profile,
):
    saved_document = Document(
        str(path)
    )

    saved_profile = (
        locate_profile_paragraph(
            saved_document
        ).text.strip()
    )

    if (
        saved_profile
        != expected_profile
    ):
        raise RuntimeError(
            "The saved CV profile did not "
            "match the approved profile."
        )


def save_manifest(
    output_path,
    verification,
    evidence,
    master_hash,
    proposed_profile,
):
    manifest = {
        "schema_version": 1,
        "created_at_utc": (
            datetime.now(
                timezone.utc
            ).isoformat()
        ),
        "applicant_name": (
            evidence.get(
                "applicant_name",
                "",
            )
        ),
        "professor_name": (
            verification.get(
                "name",
                "",
            )
        ),
        "candidate_id": (
            verification.get(
                "candidate_id",
                "",
            )
        ),
        "research_topic": (
            verification.get(
                "research_topic",
                "",
            )
        ),
        "verified_overlap_terms": (
            verification.get(
                "verified_overlap_terms",
                [],
            )
        ),
        "source_master_cv": (
            MASTER_CV.name
        ),
        "source_master_sha256": (
            master_hash
        ),
        "output_file": str(
            output_path.relative_to(
                BASE_FOLDER
            )
        ),
        "output_sha256": (
            sha256_file(
                output_path
            )
        ),
        "profile_text": (
            proposed_profile
        ),
        "master_cv_unchanged": (
            sha256_file(
                MASTER_CV
            )
            == master_hash
        ),
    }

    temporary_manifest = (
        MANIFEST_FILE.with_suffix(
            ".json.tmp"
        )
    )

    temporary_manifest.write_text(
        json.dumps(
            manifest,
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    temporary_manifest.replace(
        MANIFEST_FILE
    )


def main():
    print("=" * 68)
    print(
        "RESEARCH APPLICATION ASSISTANT — "
        "SAFE CV PROFILE UPDATE"
    )
    print("=" * 68)

    print()
    print(
        "This process does not use AI "
        "or consume API quota."
    )
    print(
        "The first three master-profile "
        "sentences remain unchanged."
    )
    print(
        "Only the final research-interest "
        "sentence can change."
    )

    temporary_path = None

    try:
        (
            verification,
            evidence,
        ) = require_approved_inputs()

        master_hash = sha256_file(
            MASTER_CV
        )

        document = Document(
            str(MASTER_CV)
        )

        profile_paragraph = (
            locate_profile_paragraph(
                document
            )
        )

        original_profile = (
            profile_paragraph.text.strip()
        )

        proposed_profile = (
            build_simple_profile(
                original_profile,
                verification,
            )
        )

        validate_profile(
            original_profile,
            proposed_profile,
        )

    except Exception as error:
        print()
        print(
            "CV tailoring could not start: "
            f"{error}"
        )
        print(
            "Master_CV.docx was not changed."
        )
        input(
            "\nPress Enter to close..."
        )
        return

    print()
    print(
        "Professor: "
        f"{verification.get('name', '')}"
    )
    print(
        "Verified topic: "
        f"{verification.get('research_topic', '')}"
    )

    print()
    print("Proposed profile:")
    print()
    print(proposed_profile)

    print()
    print(
        "Word count: "
        f"{len(proposed_profile.split())}"
    )

    print()
    print(
        "No qualifications, skills, dates, "
        "experience, or publications will "
        "be changed."
    )

    confirmation = input(
        "\nType YES to create a "
        "separate CV copy: "
    ).strip().upper()

    if confirmation != "YES":
        print(
            "Cancelled. No CV file was "
            "created or changed."
        )
        input(
            "\nPress Enter to close..."
        )
        return

    try:
        replace_paragraph_text(
            profile_paragraph,
            proposed_profile,
        )

        output_path = unique_output_path(
            evidence.get(
                "applicant_name",
                "Applicant",
            ),
            verification.get(
                "name",
                "Professor",
            ),
        )

        temporary_path = (
            output_path.with_suffix(
                ".tmp.docx"
            )
        )

        document.save(
            str(temporary_path)
        )

        if (
            sha256_file(MASTER_CV)
            != master_hash
        ):
            raise RuntimeError(
                "Master_CV.docx changed "
                "unexpectedly."
            )

        verify_saved_copy(
            temporary_path,
            proposed_profile,
        )

        temporary_path.replace(
            output_path
        )

        temporary_path = None

        save_manifest(
            output_path,
            verification,
            evidence,
            master_hash,
            proposed_profile,
        )

        if (
            sha256_file(MASTER_CV)
            != master_hash
        ):
            raise RuntimeError(
                "Master_CV.docx changed "
                "unexpectedly after saving."
            )

    except Exception as error:
        if (
            temporary_path
            and temporary_path.exists()
        ):
            temporary_path.unlink()

        print()
        print(
            "The CV copy could not be "
            f"saved safely: {error}"
        )
        print(
            "Master_CV.docx was not "
            "intentionally changed."
        )
        input(
            "\nPress Enter to close..."
        )
        return

    print()
    print(
        "Tailored CV copy created "
        "successfully."
    )
    print(
        f"Saved to: {output_path}"
    )
    print(
        "Master_CV.docx remains unchanged."
    )
    print(
        "Open the copy in Word and inspect "
        "every page before use."
    )

    input(
        "\nPress Enter to close this window..."
    )


if __name__ == "__main__":
    main()
