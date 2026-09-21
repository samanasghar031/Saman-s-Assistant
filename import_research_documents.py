import hashlib
import json
import os
import re
import shutil
from datetime import datetime, timezone
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox

from docx import Document

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None


BASE_FOLDER = Path(__file__).resolve().parent

DOCUMENT_FOLDER = (
    BASE_FOLDER / "student_documents"
)

EVIDENCE_FILE = (
    BASE_FOLDER / "student_research_evidence.json"
)

REVIEW_FILE = (
    BASE_FOLDER
    / "student_research_evidence_review.txt"
)

ARCHIVE_FOLDER = (
    BASE_FOLDER / "student_evidence_archive"
)

MAX_FILE_SIZE = 30 * 1024 * 1024
SUMMARY_SEARCH_LIMIT = 5_500
PUBLICATION_SEARCH_LIMIT = 5_500
REVIEW_TEXT_LIMIT = 40_000


def sha256_file(path):
    digest = hashlib.sha256()

    with path.open("rb") as file:
        for block in iter(
            lambda: file.read(1024 * 1024),
            b"",
        ):
            digest.update(block)

    return digest.hexdigest()


def clean_text(text):
    text = (text or "").replace(
        "\x00",
        "",
    )

    text = text.replace(
        "\r\n",
        "\n",
    ).replace(
        "\r",
        "\n",
    )

    text = re.sub(
        r"[ \t]+",
        " ",
        text,
    )

    text = re.sub(
        r"\n[ \t]+",
        "\n",
        text,
    )

    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text,
    )

    return text.strip()


def validate_source(path, extension):
    if not path.exists() or not path.is_file():
        raise FileNotFoundError(
            f"File was not found: {path}"
        )

    if path.suffix.lower() != extension:
        raise ValueError(
            f"Expected a {extension} file: "
            f"{path.name}"
        )

    if path.stat().st_size > MAX_FILE_SIZE:
        raise ValueError(
            f"{path.name} is larger than 30 MB."
        )


def extract_pdf(path):
    if PdfReader is None:
        raise RuntimeError(
            "pypdf is not installed. "
            "Run: pip install pypdf"
        )

    reader = PdfReader(
        str(path)
    )

    page_text = []

    for page_number, page in enumerate(
        reader.pages,
        start=1,
    ):
        text = clean_text(
            page.extract_text() or ""
        )

        if text:
            page_text.append(
                f"[PAGE {page_number}]\n{text}"
            )

    result = "\n\n".join(
        page_text
    ).strip()

    if len(result) < 200:
        raise ValueError(
            "The PDF contains too little "
            "extractable text. It may be a "
            "scanned image PDF."
        )

    return result, len(reader.pages)


def extract_docx(path):
    document = Document(
        str(path)
    )

    sections = []

    for paragraph in document.paragraphs:
        text = clean_text(
            paragraph.text
        )

        if text:
            sections.append(text)

    for table_number, table in enumerate(
        document.tables,
        start=1,
    ):
        rows = []

        for row in table.rows:
            cells = [
                clean_text(cell.text)
                for cell in row.cells
            ]

            if any(cells):
                rows.append(
                    " | ".join(cells)
                )

        if rows:
            sections.append(
                f"[TABLE {table_number}]\n"
                + "\n".join(rows)
            )

    result = "\n\n".join(
        sections
    ).strip()

    if len(result) < 100:
        raise ValueError(
            "The Word document contains "
            "too little readable text."
        )

    return result


def select_files():
    root = tk.Tk()
    root.withdraw()
    root.attributes(
        "-topmost",
        True,
    )

    messagebox.showinfo(
        "Import research evidence",
        "First select the applicant's "
        "published article PDF.",
        parent=root,
    )

    publication = (
        filedialog.askopenfilename(
            title="Select publication PDF",
            filetypes=[
                ("PDF files", "*.pdf")
            ],
            parent=root,
        )
    )

    if not publication:
        root.destroy()
        return None, None

    messagebox.showinfo(
        "Import research evidence",
        "Now select the thesis abstract or "
        "research-summary Word file.",
        parent=root,
    )

    summary = (
        filedialog.askopenfilename(
            title="Select research summary",
            filetypes=[
                (
                    "Word documents",
                    "*.docx",
                )
            ],
            parent=root,
        )
    )

    root.destroy()

    if not summary:
        return None, None

    return (
        Path(publication),
        Path(summary),
    )


def build_review(
    applicant_name,
    publication_path,
    publication_text,
    publication_pages,
    summary_path,
    summary_text,
):
    return f"""STUDENT RESEARCH EVIDENCE — REVIEW BEFORE APPROVAL
{"=" * 72}

Applicant: {applicant_name}
Created: {datetime.now(timezone.utc).isoformat()}

IMPORTANT

Only approve if the selected files belong to this applicant and the
extracted text is readable. Approval does not verify claims from any
other CV, website, professor page, or AI response.

PUBLICATION SOURCE
{"-" * 72}

Filename: {publication_path.name}
Pages: {publication_pages}
SHA-256: {sha256_file(publication_path)}
Extracted characters: {len(publication_text)}

{publication_text[:REVIEW_TEXT_LIMIT]}

RESEARCH SUMMARY OR THESIS ABSTRACT SOURCE
{"-" * 72}

Filename: {summary_path.name}
SHA-256: {sha256_file(summary_path)}
Extracted characters: {len(summary_text)}

{summary_text[:REVIEW_TEXT_LIMIT]}

{"=" * 72}
END OF REVIEW
"""


def safe_name(text):
    value = re.sub(
        r"[^A-Za-z0-9._-]+",
        "_",
        text,
    ).strip("._")

    return value or "document"


def archive_previous_evidence():
    if not EVIDENCE_FILE.exists():
        return None

    ARCHIVE_FOLDER.mkdir(
        exist_ok=True
    )

    stamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    destination = (
        ARCHIVE_FOLDER
        / f"student_research_evidence_{stamp}.json"
    )

    shutil.copy2(
        EVIDENCE_FILE,
        destination,
    )

    return destination


def save_approved_evidence(
    applicant_name,
    publication_path,
    publication_text,
    publication_pages,
    summary_path,
    summary_text,
):
    stamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    import_folder = (
        DOCUMENT_FOLDER / stamp
    )

    import_folder.mkdir(
        parents=True,
        exist_ok=False,
    )

    publication_copy = (
        import_folder
        / (
            "publication_"
            + safe_name(
                publication_path.name
            )
        )
    )

    summary_copy = (
        import_folder
        / (
            "research_summary_"
            + safe_name(
                summary_path.name
            )
        )
    )

    publication_text_file = (
        import_folder
        / "publication_extracted.txt"
    )

    summary_text_file = (
        import_folder
        / "research_summary_extracted.txt"
    )

    shutil.copy2(
        publication_path,
        publication_copy,
    )

    shutil.copy2(
        summary_path,
        summary_copy,
    )

    publication_text_file.write_text(
        publication_text,
        encoding="utf-8",
    )

    summary_text_file.write_text(
        summary_text,
        encoding="utf-8",
    )

    approved_search_evidence = (
        "RESEARCH SUMMARY OR "
        "THESIS ABSTRACT:\n"
        + summary_text[
            :SUMMARY_SEARCH_LIMIT
        ]
        + "\n\nPUBLICATION TEXT:\n"
        + publication_text[
            :PUBLICATION_SEARCH_LIMIT
        ]
    )

    evidence = {
        "approved_search_evidence": (
            approved_search_evidence
        ),
        "schema_version": 1,
        "applicant_name": applicant_name,
        "approved_by_user": True,
        "approved_at_utc": (
            datetime.now(
                timezone.utc
            ).isoformat()
        ),
        "usage_rules": [
            (
                "Use only facts present in "
                "the approved source text."
            ),
            (
                "Do not infer unlisted skills, "
                "results, authorship, or experience."
            ),
            (
                "Keep separate research projects "
                "separate unless both are relevant."
            ),
            (
                "Professor recruitment and "
                "publications require independent "
                "checks."
            ),
        ],
        "sources": [
            {
                "type": "publication",
                "original_filename": (
                    publication_path.name
                ),
                "stored_file": str(
                    publication_copy.relative_to(
                        BASE_FOLDER
                    )
                ),
                "extracted_text_file": str(
                    publication_text_file.relative_to(
                        BASE_FOLDER
                    )
                ),
                "sha256": sha256_file(
                    publication_path
                ),
                "pages": publication_pages,
            },
            {
                "type": "research_summary",
                "original_filename": (
                    summary_path.name
                ),
                "stored_file": str(
                    summary_copy.relative_to(
                        BASE_FOLDER
                    )
                ),
                "extracted_text_file": str(
                    summary_text_file.relative_to(
                        BASE_FOLDER
                    )
                ),
                "sha256": sha256_file(
                    summary_path
                ),
            },
        ],
    }

    temporary = EVIDENCE_FILE.with_suffix(
        ".json.tmp"
    )

    temporary.write_text(
        json.dumps(
            evidence,
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    temporary.replace(
        EVIDENCE_FILE
    )

    return import_folder


def open_review_file():
    try:
        os.startfile(
            str(REVIEW_FILE)
        )

    except AttributeError:
        print(
            "Open this file manually: "
            f"{REVIEW_FILE}"
        )

    except OSError as error:
        print(
            "Windows could not open the "
            f"review file: {error}"
        )


def main():
    print("=" * 68)
    print(
        "RESEARCH APPLICATION ASSISTANT — "
        "IMPORT RESEARCH EVIDENCE"
    )
    print("=" * 68)

    print()
    print(
        "This process does not use Gemini "
        "or consume API quota."
    )
    print(
        "It extracts text locally and waits "
        "for your approval."
    )

    applicant_name = input(
        "\nApplicant's full name: "
    ).strip()

    if not applicant_name:
        print(
            "Applicant name is required."
        )
        input(
            "\nPress Enter to close..."
        )
        return

    (
        publication_path,
        summary_path,
    ) = select_files()

    if (
        not publication_path
        or not summary_path
    ):
        print(
            "File selection was cancelled. "
            "No evidence was changed."
        )
        input(
            "\nPress Enter to close..."
        )
        return

    try:
        validate_source(
            publication_path,
            ".pdf",
        )

        validate_source(
            summary_path,
            ".docx",
        )

        print()
        print(
            "Step 1 of 4: Extracting "
            "the publication PDF..."
        )

        (
            publication_text,
            publication_pages,
        ) = extract_pdf(
            publication_path
        )

        print(
            "Step 2 of 4: Extracting "
            "the research summary..."
        )

        summary_text = extract_docx(
            summary_path
        )

        print(
            "Step 3 of 4: Creating "
            "the review file..."
        )

        review = build_review(
            applicant_name,
            publication_path,
            publication_text,
            publication_pages,
            summary_path,
            summary_text,
        )

        REVIEW_FILE.write_text(
            review,
            encoding="utf-8",
        )

        open_review_file()

    except Exception as error:
        print()
        print(
            "Evidence import could not "
            f"be prepared: {error}"
        )
        print(
            "The previous approved evidence "
            "was not changed."
        )
        input(
            "\nPress Enter to close..."
        )
        return

    print()
    print(
        f"Review opened: {REVIEW_FILE.name}"
    )
    print(
        "Check the applicant name, "
        "filenames, and extracted text."
    )

    approval = input(
        "\nType APPROVE only if the "
        "extraction is correct: "
    ).strip().upper()

    if approval != "APPROVE":
        print(
            "Not approved. Existing approved "
            "evidence was not changed."
        )
        input(
            "\nPress Enter to close..."
        )
        return

    try:
        archived = archive_previous_evidence()

        import_folder = save_approved_evidence(
            applicant_name,
            publication_path,
            publication_text,
            publication_pages,
            summary_path,
            summary_text,
        )

    except Exception as error:
        print()
        print(
            "Approved evidence could not "
            f"be saved: {error}"
        )
        input(
            "\nPress Enter to close..."
        )
        return

    print()
    print(
        "Step 4 of 4: Approved "
        "evidence saved."
    )
    print(
        f"Evidence file: {EVIDENCE_FILE.name}"
    )
    print(
        f"Document copies: {import_folder}"
    )

    if archived:
        print(
            "Previous evidence backup: "
            f"{archived}"
        )

    print(
        "Future search and verification "
        "stages can now use these sources."
    )

    input(
        "\nPress Enter to close this window..."
    )


if __name__ == "__main__":
    main()
